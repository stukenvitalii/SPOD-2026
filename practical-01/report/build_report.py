from pathlib import Path
from collections import defaultdict
from statistics import median
import csv, math, shutil, json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION_START, WD_ORIENT

ROOT = Path(__file__).resolve().parent
LAB = ROOT.parent
A = ROOT / 'assets'
A.mkdir(exist_ok=True)
source = ROOT / 'measurements.csv'
rows = list(csv.DictReader(source.open(encoding='utf-8-sig')))
assert len(rows) == 180, f'Expected 180 measurements, got {len(rows)}'
groups = defaultdict(list)
for row in rows:
    groups[row['mode'], int(row['processes']), int(row['count'])].append(float(row['seconds_per_round']))
times = {key:median(values) for key,values in groups.items()}
ps = [1,2,4,8,16,32]
counts = [1,256,16384]

def font(n): return ImageFont.truetype('C:/Windows/Fonts/arial.ttf', n)

def chart(mode, speed=False):
    im=Image.new('RGB',(1600,760),'white'); d=ImageDraw.Draw(im)
    l,r,t,b=170,1500,100,620
    vals=[[times[mode,1,c]/times[mode,p,c] if speed else times[mode,p,c]*1e6 for p in ps] for c in counts]
    low=math.floor(math.log10(min(min(v) for v in vals))); high=math.ceil(math.log10(max(max(v) for v in vals)))
    if high==low: high+=1
    def x(i):return l+(r-l)*i/5
    def y(v):return b-(math.log10(v)-low)/(high-low)*(b-t)
    d.text((800,25),('Ускорение S = T1 / Tp' if speed else 'Время одного раунда в микросекундах')+'   '+mode,font=font(32),fill='black',anchor='mt')
    for exp in range(low,high+1):
        yy=y(10**exp); d.line((l,yy,r,yy),fill='#d8d8d8',width=2)
        d.text((l-15,yy),f'{10**exp:g}',font=font(23),fill='black',anchor='rm')
    d.line((l,t,l,b,r,b),fill='black',width=3)
    for i,p in enumerate(ps):d.text((x(i),b+18),str(p),font=font(25),fill='black',anchor='mt')
    for c,values,color in zip(counts,vals,['#245d9c','#c05c17','#218352']):
        points=[(x(i),y(v)) for i,v in enumerate(values)]
        d.line(points,fill=color,width=5)
        for xx,yy in points:d.ellipse((xx-6,yy-6,xx+6,yy+6),fill=color)
        j=counts.index(c); xx=220+j*440
        d.line((xx,715,xx+50,715),fill=color,width=5)
        d.text((xx+65,715),f'n = {c} ({4*c} байт)',font=font(24),fill='black',anchor='lm')
    d.text((800,675),'Число процессов p     Шкала Y логарифмическая',font=font(23),fill='black',anchor='mm')
    path=A/f'{mode}-{"speed" if speed else "time"}.png'; im.save(path); return path

for m in ['unordered','ordered']:
    chart(m);chart(m,True)
im=Image.open(A/'runs.png'); im.crop((342,35,1400,385)).save(A/'runs-crop.png')

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(1.5)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(14)
normal.paragraph_format.line_spacing=1;normal.paragraph_format.space_after=Pt(5)
for name,size in [('Title',14),('Heading 1',14),('Heading 2',14)]:
    s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(9)
doc.styles['Subtitle'].font.color.rgb=RGBColor(0,0,0)
for st in doc.styles:
    pr=st.element.find(qn('w:pPr'))
    if pr is not None:
        for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
footer=sec.footer.paragraphs[0];footer.alignment=1
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
pg=OxmlElement('w:pgNumType');pg.set(qn('w:start'),'2');sec._sectPr.append(pg)

def p(s):doc.add_paragraph(s)
def h(s):doc.add_heading(s,1)
def page(s):h(s)
def picture(path,width=17):
    gray=A/(Path(path).stem+'-gray.png')
    Image.open(path).convert('L').save(gray)
    par=doc.add_paragraph();par.alignment=1;par.paragraph_format.keep_with_next=True
    selected = path if Path(path).stem.endswith(('-time', '-speed')) else gray
    par.add_run().add_picture(str(selected),width=Cm(min(width,17)))
def caption(s):
    par=doc.add_paragraph(s);par.alignment=1
    for run in par.runs:run.font.size=Pt(14)
def table(headers,data):
    tb=doc.add_table(rows=1,cols=len(headers));tb.style='Table Grid'
    for cell,txt in zip(tb.rows[0].cells,headers):cell.text=str(txt)
    pr=tb.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');pr.append(repeat)
    for row in data:
        for cell,txt in zip(tb.add_row().cells,row):cell.text=str(txt)
    for row in tb.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(3);par.paragraph_format.space_before=Pt(3)
                for run in par.runs:run.font.size=Pt(14)
    for cell in tb.rows[0].cells:
        for run in cell.paragraphs[0].runs:run.bold=True
    doc.add_paragraph()

doc.add_paragraph('Практическая работа 1',style='Title')
doc.add_paragraph('Обмен сообщениями между процессами средствами MPI',style='Subtitle')
h('1 Вариант и постановка задачи')
p('Вариант — общее задание для всех студентов. Цель работы — реализовать обмен рангами между MPI-процессами, исследовать порядок получения сообщений и сравнить время двух способов организации приёма.')
p('Задание 1. Каждый процесс с ненулевым рангом отправляет процессу 0 значение своего ранга. Процесс 0 сначала печатает собственный ранг, затем принимает p − 1 сообщений от произвольных источников и после каждого приёма печатает полученное значение. Выполнить повторные запуски и объяснить порядок вывода.')
p('Задание 2. Изменить программу так, чтобы ранги выводились по возрастанию: 0, 1, …, p − 1. Сравнить результаты и время обеих версий, представить сети Петри и исследовать зависимости от числа процессов и размера сообщения.')
h('2 Алгоритмы решения')
p('Каждый процесс выполняет MPI_Init, узнаёт rank и size в MPI_COMM_WORLD и создаёт собственный массив целых чисел. Все элементы массива отправителя равны его рангу. По умолчанию массив содержит один элемент; увеличение n используется для исследования объёма обмена.')
p('В unordered.cpp нулевой процесс принимает с MPI_ANY_SOURCE. Порядок между разными отправителями не фиксируется; печать следует за фактическим приёмом. В ordered.cpp источники последовательно задаются от 1 до size − 1. Доступное сообщение от большего ранга не заменяет ожидаемое сообщение от меньшего.')
p('Количество приёмов определяется size, а не константой. При p = 1 нулевой процесс только печатает свой ранг. При повторении раундов тег равен номеру раунда, поэтому сообщения соседних раундов не смешиваются. Используются стандартные блокирующие MPI_Send и MPI_Recv; завершение Send не означает завершения печати получателем.')

page('3 Условия и методика эксперимента')
table(['Параметр','Значение'],[
('Дата измерений','28 октября 2026 года'),('Процессор','AMD Ryzen 7 8845H with Radeon 780M Graphics'),
('Доступные логические процессоры','16'),('Среда','Windows x64, локальный запуск на одном компьютере'),
('Компилятор','GCC 14.2.0, MSYS2 UCRT64'),('Параметры сборки','C++17, -O2 -Wall -Wextra -Wpedantic'),
('Реализация MPI','Microsoft MPI Runtime 10.1.12498.52'),('Число процессов p','1, 2, 4, 8, 16, 32'),
('Размер n','1, 256, 16384 элементов MPI_INT'),('Размер пакета','4, 1024, 65536 байт'),
('Серия','1000 раундов; 5 независимых повторений'),('Прогрев','10 раундов отдельным запуском перед серией')])
p('Измерения выполнены режимом --benchmark без приветствий. После подготовки данных выполняется MPI_Barrier, затем каждый процесс измеряет серию через MPI_Wtime. Максимум локальных интервалов собирается MPI_Reduce с операцией MPI_MAX. Подготовка данных, исходный барьер, итоговая редукция, консольный вывод, запуск процессов и MPI_Finalize в интервал не входят.')
p('Для каждого повтора время серии делится на 1000. В таблицах приведена медиана пяти значений времени раунда. Между раундами нет барьера: измеряется серия с возможным перекрытием работы участников, а не изолированная задержка одного сообщения. Начальный барьер не означает абсолютно одновременный старт часов.')
p('При p = 32 число процессов превышает число доступных логических процессоров. Привязка процессов к ядрам специально не задавалась. Фоновая нагрузка ОС не исключалась; небольшие немонотонности не следует объяснять единственной причиной без дополнительного профилирования.')

for mode,title in [('unordered','произвольного приёма'),('ordered','упорядоченного приёма')]:
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.PORTRAIT;sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    h('4 Формальная модель '+title)
    p('Обычный режим, три процесса, один раунд. Начальная маркировка M₀ = {start, s1, s2}; во всех остальных позициях маркеров нет. Все дуги имеют вес 1.')
    picture(LAB/'diagrams'/f'{mode}.png',17)
    caption('Рисунок '+('1' if mode=='unordered' else '2')+' — Сеть Петри '+title)
sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.PORTRAIT;sec.page_width=Cm(21);sec.page_height=Cm(29.7)
h('4 Продолжение описания моделей')
p('Send обозначает логическую доступность отдельного сообщения, а не точный момент возврата MPI_Send. Внутренние буферы MPI и MPI_Finalize скрыты. Сообщения m1 и m2 различаются по отправителю; done означает окончание требуемой печати.')
table(['Действие','Функция перехода'],[('Send1, Send2','Публикация сообщения соответствующего ранга'),('Print0','Печать нуля до первого приёма'),('Recv1, Recv2','Потребление сообщения и переход в состояние принятого ранга'),('Print1, Print2','Печать принятого ранга и разрешение продолжить'),('End в unordered','Проверка завершения двух пар приём — печать')])
p('В произвольной версии Recv1 и Recv2 конкурируют за единственный маркер ready. Приём требует также m1 или m2. После приёма ready отсутствует до завершения печати. Print1 и Print2 возвращают ready и создают f1 или f2; End требует все три условия ready, f1, f2. Позиции f1 и f2 формально представляют выполненную часть цикла, а не дополнительные переменные C++.')
p('В упорядоченной версии q1 разрешает только приём от 1. После его печати появляется q2, разрешающий приём от 2. Отправки могут происходить до Print0, однако приём до печати нуля запрещён обеими сетями.')
table(['Свойство','unordered','ordered'],[('Достижимых маркировок','20','14'),('Максимум маркеров в позиции','1','1'),('Нештатных тупиков','0','0'),('Завершение достижимо из каждого состояния','Да','Да'),('Последовательности печати','0,1,2 и 0,2,1','0,1,2')])
p('Результат получен полным перебором достижимых состояний конечных моделей. Терминальная маркировка {done} является штатным завершением. Проверка относится к трём процессам и одному раунду; она не подменяет доказательство всей MPI-реализации для произвольных параметров.')

page('5 Примеры работы и проверка результата')
p('На рисунке 3 показан реальный вывод двух последовательных запусков на восьми процессах. Первые восемь строк относятся к unordered, следующие восемь — к ordered.')
picture(A/'runs-crop.png')
caption('Рисунок 3 — Вывод обеих версий в терминале Visual Studio Code')
p('В первом запуске получен порядок 0, 7, 4, 6, 2, 5, 1, 3. Во втором — 0, 1, 2, 3, 4, 5, 6, 7. Оба результата соответствуют заданию: все ранги присутствуют по одному разу, ноль напечатан первым. В первой версии порядок определяется сопоставлением сообщений с приёмами MPI_ANY_SOURCE и может изменяться; повторение одинакового порядка не является ошибкой.')
table(['Повтор unordered при p = 8','Последовательность рангов'],[[x['run'],x['ranks']] for x in json.loads((ROOT/'examples.json').read_text(encoding='utf-8-sig'))])
p('Автоматическая проверка check.ps1 выполняет обе версии на 1, 2, 4, 8 и 16 процессах по три раунда. Контролируются число строк, ноль первым, полный набор рангов без повторов и возрастающий порядок ordered. Отдельно проверяется формат измерений при сообщениях по 64 КиБ. При p = 1 вывод состоит только из Hello from 0.')
p('Блокирующая отправка не создаёт циклического ожидания: процесс 0 не выполняет встречную отправку перед приёмами и последовательно обслуживает всех отправителей. Корректность не требует предполагать бесконечную внутреннюю буферизацию MPI.')

page('6 Таблицы времени выполнения')
p('Единица измерения — микросекунда на раунд. Каждая ячейка содержит медиану пяти независимых серий. n — число целых элементов в одном сообщении. В раунде передаётся p − 1 сообщений, то есть 4n(p − 1) байт полезных данных.')
for mode in ['unordered','ordered']:
    doc.add_heading(mode,2)
    table(['p']+[f'n = {c}' for c in counts],[[p0]+[f'{times[mode,p0,c]*1e6:.4f}' for c in counts] for p0 in ps])
p('При p = 1 сообщение не пересылается независимо от n. Измеряется только остаточная работа измеряемого участка и часов. Эти малые значения чувствительны к разрешению таймера и оптимизации компилятора. Они сохранены в таблице для полноты, но не характеризуют скорость передачи данных.')

page('7 Графики времени выполнения')
for i,mode in enumerate(['unordered','ordered']):
    picture(A/f'{mode}-time.png');caption(f'Рисунок {4+i} — Время раунда для {mode}')
p('По оси X отложено p, разные кривые соответствуют размерам n. Логарифмическая шкала Y позволяет одновременно показать вырожденный случай p = 1 и обмены при p ≥ 2. Соединяющие отрезки служат для чтения дискретных экспериментальных точек.')

page('8 Формальное ускорение и границы интерпретации')
p('По требованию исследования вычислено Sₚ(n) = T₁(n) / Tₚ(n), где T — медианное время раунда. Значения ниже рассчитаны именно по этой формуле. При p = 1 значение S равно 1 по определению.')
for mode in ['unordered','ordered']:
    doc.add_heading(mode,2)
    table(['p']+[f'n = {c}' for c in counts],[[p0]+[f'{times[mode,1,c]/times[mode,p0,c]:.3e}' for c in counts] for p0 in ps])
p('Для данной задачи T₁ не является временем последовательного выполнения того же объёма обменов: при одном процессе отправителей нет. Кроме того, при увеличении p возрастает число сообщений. Поэтому Sₚ представляет формальное отношение времён разных коммуникационных нагрузок; его нельзя интерпретировать как ускорение фиксированной вычислительной задачи.')
p('Значения Sₚ < 1 показывают рост измеренного времени относительно вырожденной базы p = 1, но сами по себе не доказывают неэффективность MPI. По этой постановке нельзя установить область ускорения одинаковой полезной работы. Сравнение ordered и unordered корректнее выполнять при одинаковых p и n.')

page('9 Графики формального ускорения')
for i,mode in enumerate(['unordered','ordered']):
    picture(A/f'{mode}-speed.png');caption(f'Рисунок {6+i} — Формальное отношение T₁/Tₚ для {mode}')
p('Шкала Y логарифмическая. Резкий переход от p = 1 к p = 2 связан прежде всего с появлением обмена, отсутствующего в базе, и не является точкой потери ускорения фиксированной задачи.')

page('10 Анализ результатов и выводы')
for mode in ['unordered','ordered']:
    t2=times[mode,2,16384]*1e6;t32=times[mode,32,16384]*1e6
    p(f'Для {mode} при пакете 64 КиБ медианное время изменяется от {t2:.2f} мкс при p = 2 до {t32:.2f} мкс при p = 32. Отношение этих времён составляет {t32/t2:.2f}. Это сравнение серий с разным числом сообщений, а не оценка ускорения неизменной нагрузки.')
for c in counts:
    ratio=times['ordered',8,c]/times['unordered',8,c]
    p(f'При p = 8 и n = {c} отношение времени ordered к unordered равно {ratio:.3f}. Значение больше 1 означает более медленный ordered в данном эксперименте, меньше 1 — более быстрый. Порядок обработки сам по себе не гарантирует универсального преимущества одной версии.')
p('Рост времени объясняется увеличением числа сообщений, последовательной обработкой у процесса 0 и накладными расходами MPI. При крупных пакетах добавляется стоимость копирования и передачи данных. Участок p = 32 дополнительно связан с превышением числа доступных логических процессоров. По одним временным измерениям нельзя выделить вклад планировщика или смены протокола обмена.')
p('Небольшие подъёмы и спады между соседними точками могут отражать планирование, буферизацию и фоновую нагрузку. Локальное снижение времени не следует объявлять устойчивой зоной ускорения без проверки разброса и сопоставимости работы. Медиана уменьшает влияние отдельных выбросов, но не устраняет систематические различия условий запуска.')
p('Обе реализации решают поставленную задачу. При произвольном источнике порядок рангов недетерминирован; при фиксированных источниках порядок гарантируется цепочкой приёмов. Сети Петри сохраняют источник сообщения и зависимость печати от приёма; проверенные конечные модели безопасны и не содержат нештатных тупиков. Для работы освоены парные обмены MPI, параметризация числа процессов и измерение времени без ввода и вывода.')
p('Область полезного применения этой программы — демонстрация и исследование коммуникации. Обязательное вычислительное ускорение для неё не следует из постановки. Для доказательства ускорения понадобилась бы задача с фиксированным объёмом полезных вычислений и содержательной последовательной базой.')

page('Приложение А Исходный код')
p('Листинги соответствуют программам, использованным в эксперименте. common.hpp содержит общие параметры, проверку аргументов, печать и завершение измерения. Два файла cpp собираются в отдельные исполняемые программы.')
for name in ['unordered.cpp','ordered.cpp','common.hpp']:
    doc.add_heading(name,2)
    for line in (LAB/'src'/name).read_text(encoding='utf-8').splitlines():
        par=doc.add_paragraph();par.paragraph_format.space_after=Pt(0);par.paragraph_format.line_spacing=1
        run=par.add_run(line or ' ');run.font.name='Consolas';run.font.size=Pt(8)

page('Приложение Б Воспроизведение эксперимента')
p('Из корня проекта в PowerShell:')
for command in ['.\\practical-01\\build.ps1','.\\practical-01\\check.ps1','.\\practical-01\\run.ps1 -Mode unordered -Processes 8','.\\practical-01\\run.ps1 -Mode ordered -Processes 8','.\\practical-01\\benchmark.ps1 -Processes 1,2,4,8,16,32 -Counts 1,256,16384 -Rounds 1000 -Repeats 5']:
    par=doc.add_paragraph(command);par.runs[0].font.name='Consolas';par.runs[0].font.size=Pt(9)
p('Первичные данные сохранены рядом с отчётом в measurements.csv: режим, число процессов, размер пакета в элементах, число раундов, время серии в секундах, время раунда и номер повтора. Для каждого сочетания параметров выполнено пять измерений. Полные исходники и скрипты находятся в practical-01.')
p('На macOS используются те же исходники и Open MPI: bash practical-01/build.sh, затем bash practical-01/run.sh unordered 8 или ordered 8. Представленные в отчёте измерения относятся к Windows и Microsoft MPI; результаты другой машины следует исследовать отдельной серией.')
h('Использованные материалы')
p('1. Содержание отчёта по практической работе СПОД. Требования преподавателя, предоставленные в Markdown.')
p('2. Конспекты лекций 3 и 4 по дисциплине «Системы параллельной обработки данных»: постановка первой работы, MPI и требования к исследованию.')
p('3. Мараховский В. Б., Розенблюм Л. Я., Яковлев А. В. Моделирование параллельных процессов. Сети Петри. СПб.: Профессиональная литература, 2014. Литература, указанная в требованиях к формальному описанию.')
doc.core_properties.title='Практическая работа 1 Обмен сообщениями средствами MPI'
doc.core_properties.author=''
for style in doc.styles:
    if hasattr(style,'font'):
        style.font.name='Times New Roman';style.font.size=Pt(14);style.font.color.rgb=RGBColor(0,0,0)
        rp=style.element.find(qn('w:rPr'))
        if rp is not None:
            rf=rp.find(qn('w:rFonts'))
            if rf is not None:
                for attribute in list(rf.attrib):
                    if 'Theme' in attribute:del rf.attrib[attribute]
for run_element in doc.element.iter(qn('w:r')):
    rp=run_element.find(qn('w:rPr'))
    if rp is None:rp=OxmlElement('w:rPr');run_element.insert(0,rp)
    for tag in ['rFonts','sz','szCs','color']:
        for old in list(rp.findall(qn('w:'+tag))):rp.remove(old)
    rf=OxmlElement('w:rFonts')
    for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Times New Roman')
    rp.append(rf)
    for tag,value in [('sz','28'),('szCs','28'),('color','000000')]:
        element=OxmlElement('w:'+tag);element.set(qn('w:val'),value);rp.append(element)
for section in list(doc.sections)[1:]:
    for numbering in list(section._sectPr.findall(qn('w:pgNumType'))):section._sectPr.remove(numbering)
doc.save(ROOT/'Практическая работа 1 цветные графики.docx')
print(json.dumps({'rows':len(rows),'file':str(ROOT/'Практическая работа 1 цветные графики.docx')},ensure_ascii=True,indent=2))
